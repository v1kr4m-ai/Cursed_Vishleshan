"""
Export a summary (and chapters) to Word (.docx) and PDF.

GUI-free. Word files are built with python-docx; PDFs are printed from an HTML copy by Microsoft Edge
(it is on every Windows 10/11 PC), so Hindi / Urdu / Arabic and other scripts come out shaped correctly.
"""
import html
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from markdown_it import MarkdownIt

_RTL = re.compile("[֐-ࣿיִ-﷿ﹰ-﻿]")
FORMATS = ("none", "docx", "pdf", "both")


def _split_footer(md):
    """The summary files end with '---' and an italic line (languages, backend). Split it off."""
    m = re.search(r"\n-{3,}\s*\n_([^\n]+)_\s*$", md)
    return (md[:m.start()], m.group(1).strip()) if m else (md, "")


def _title_of(path):
    stem = Path(path).stem
    stem = re.sub(r"_summary(_[^_]+)?$", "", stem)
    return stem.replace("_", " ").strip() or stem


def _is_rtl(text):
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and sum(bool(_RTL.match(c)) for c in letters) > len(letters) / 2


# ---------------------------------------------------------------- HTML + PDF
_CSS = """
body { font-family: "Segoe UI", "Nirmala UI", "Noto Sans", Arial, sans-serif; color: #16181d; margin: 0; line-height: 1.6; font-size: 11.5pt; }
h1 { font-size: 22pt; margin: 0 0 4px; } h2 { font-size: 15pt; margin: 22px 0 6px; border-bottom: 1px solid #ddd; padding-bottom: 3px; }
h3 { font-size: 12.5pt; margin: 16px 0 4px; } p, li, h1, h2, h3 { unicode-bidi: plaintext; text-align: start; }
.meta { color: #6b7280; font-size: 9.5pt; margin-bottom: 18px; } ul, ol { padding-inline-start: 22px; }
hr { border: 0; border-top: 1px solid #ddd; margin: 20px 0; } code { background: #f3f4f6; padding: 1px 4px; border-radius: 3px; }
@page { margin: 20mm 18mm; }
"""


def build_html(title, meta, md_parts):
    md = MarkdownIt(options_update={"breaks": True})
    body = "\n".join(md.render(part) for part in md_parts if part.strip())
    # dir="auto": each list / paragraph / heading takes its own direction (Urdu, Arabic -> right to left)
    body = re.sub(r"<(p|ul|ol|li|h[1-6])>", r'<\1 dir="auto">', body)
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{_CSS}</style></head><body><h1>{html.escape(title)}</h1>"
            f"<div class='meta'>{html.escape(meta)}</div>{body}</body></html>")


def _edge():
    for var in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
        base = os.environ.get(var)
        if base:
            p = Path(base) / "Microsoft" / "Edge" / "Application" / "msedge.exe"
            if p.exists():
                return str(p)
    return shutil.which("msedge")


def html_to_pdf(html_text, pdf_path):
    edge = _edge()
    if not edge:
        raise RuntimeError("Microsoft Edge was not found - needed to make the PDF.")
    tmp = Path(tempfile.mkdtemp(prefix="vidsum_pdf_"))
    try:
        src = tmp / "doc.html"
        src.write_text(html_text, encoding="utf-8")
        out = tmp / "doc.pdf"
        subprocess.run([edge, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--user-data-dir={tmp / 'profile'}", f"--print-to-pdf={out}", src.as_uri()],
                       capture_output=True, timeout=90)
        if not out.exists():
            raise RuntimeError("Edge did not produce a PDF.")
        shutil.copyfile(out, pdf_path)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- Word
def _add_runs(par, inline_tokens, rtl):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    bold = italic = False
    for t in inline_tokens:
        if t.type == "strong_open":
            bold = True
        elif t.type == "strong_close":
            bold = False
        elif t.type == "em_open":
            italic = True
        elif t.type == "em_close":
            italic = False
        elif t.type in ("text", "code_inline", "softbreak", "hardbreak"):
            text = {"softbreak": chr(10), "hardbreak": chr(10)}.get(t.type, t.content)
            run = par.add_run(text)
            run.bold, run.italic = bold or None, italic or None
            if t.type == "code_inline":
                run.font.name = "Consolas"
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.append(fonts)
            fonts.set(qn("w:cs"), "Arial")           # complex scripts (Hindi, Arabic...) fall back sensibly
            if rtl:
                rpr.append(OxmlElement("w:rtl"))


def build_docx(title, meta, md_parts, path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    if meta:
        p = doc.add_paragraph()
        r = p.add_run(meta)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    md = MarkdownIt()
    for part in md_parts:
        toks = md.parse(part)
        i, list_kind = 0, None
        while i < len(toks):
            t = toks[i]
            if t.type == "heading_open":
                inline = toks[i + 1]
                level = min(int(t.tag[1]), 3)
                h = doc.add_heading("", level)
                _add_runs(h, inline.children or [], _is_rtl(inline.content))
                if _is_rtl(inline.content):
                    h._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                i += 3
                continue
            if t.type in ("bullet_list_open", "ordered_list_open"):
                list_kind = "List Bullet" if t.type == "bullet_list_open" else "List Number"
            elif t.type in ("bullet_list_close", "ordered_list_close"):
                list_kind = None
            elif t.type == "paragraph_open":
                inline = toks[i + 1]
                p = doc.add_paragraph(style=list_kind) if list_kind else doc.add_paragraph()
                rtl = _is_rtl(inline.content)
                _add_runs(p, inline.children or [], rtl)
                if rtl:
                    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                i += 3
                continue
            elif t.type == "hr":
                doc.add_paragraph("_" * 40)
            i += 1
    doc.save(str(path))


# ---------------------------------------------------------------- public
def export_summary(md_path, fmts, chapters_path=None, out_dir=None):
    """Export one summary .md as .docx and/or .pdf next to it (or in out_dir).
    chapters_path: optional chapters .md appended as a section. Returns the files written."""
    if fmts in (None, "", "none"):
        return []
    md_path = Path(md_path)
    body, meta = _split_footer(md_path.read_text(encoding="utf-8", errors="replace"))
    parts = [body]
    if chapters_path and Path(chapters_path).exists():
        ch, _ = _split_footer(Path(chapters_path).read_text(encoding="utf-8", errors="replace"))
        ch = re.sub(r"^# [^\n]*\n", "", ch.strip() + "\n")        # drop its own title line
        parts.append("# Chapters and key points\n\n" + re.sub(r"^## ", "## ", ch))
    title = _title_of(md_path)
    dest = Path(out_dir) if out_dir else md_path.parent
    stem = dest / md_path.stem
    out = []
    if fmts in ("docx", "both"):
        build_docx(title, meta, parts, f"{stem}.docx")
        out.append(Path(f"{stem}.docx"))
    if fmts in ("pdf", "both"):
        html_to_pdf(build_html(title, meta, parts), f"{stem}.pdf")
        out.append(Path(f"{stem}.pdf"))
    return out
